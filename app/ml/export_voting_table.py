"""
Export Tabel Voting Mayoritas per Pohon Keputusan.

Kolom output:
  ID | Gejala Kerusakan | P1 ... Pn | Hasil Voting | Aktual | Evaluasi

- "Gejala Kerusakan" = daftar gejala YA dipisah ", "
- Px = prediksi pohon ke-x (singkatan label)
- Hasil Voting = majority vote dari SEMUA pohon
- --data train (default) atau --data test

Penggunaan:
    python -m app.ml.export_voting_table --max-trees 10
    python -m app.ml.export_voting_table --max-trees 10 --data test
"""

import os
import sys
import argparse
import joblib
import numpy as np
import pandas as pd
from collections import Counter
from sklearn.model_selection import train_test_split

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from app.ml.preprocessing import (
    load_dataset,
    transform_features,
    FEATURE_COLS,
    TARGET_COL,
)

DATASET_PATH  = os.path.join(BASE_DIR, "dataset", "dataset.xlsx")
MODEL_DIR     = os.path.join(os.path.dirname(__file__), "models")
MODEL_PATH    = os.path.join(MODEL_DIR, "rf_model.pkl")
ENCODERS_PATH = os.path.join(MODEL_DIR, "label_encoders.pkl")
OUTPUT_PATH   = os.path.join(MODEL_DIR, "voting_table.xlsx")

FEAT_DISPLAY = {
    "suara_bising_abnormal":             "Suara Bising Abnormal",
    "bau_hangus":                        "Bau Hangus",
    "indikasi_overheating":              "Indikasi Overheating",
    "putaran_poros_seret":               "Putaran Poros Seret",
    "getaran_berlebih":                  "Getaran Berlebih",
    "terminal_overheating":              "Terminal Overheating",
    "kipas_pendingin_rusak":             "Kipas Pendingin Rusak",
    "cooling_duct_tersumbat":            "Cooling Duct Tersumbat",
    "resistansi_isolasi_tidak_seimbang": "Res. Isolasi Tidak Seimbang",
    "resistansi_winding_tidak_seimbang": "Res. Winding Tidak Seimbang",
    "arus_antar_fasa_tidak_seimbang":    "Arus Antar Fasa Tidak Seimbang",
}

LABEL_ABBREV = {
    "Kerusakan Bearing":         "Bearing",
    "Kerusakan Gulungan Stator": "Gulungan",
    "Kerusakan Housing Bearing": "Housing",
    "Kerusakan Shaft":           "Shaft",
    "Kerusakan Terminal":        "Terminal",
}


def _abbrev(label):
    return LABEL_ABBREV.get(label, label)


def _gejala_aktif(x_row):
    aktif = [
        FEAT_DISPLAY.get(feat, feat)
        for j, feat in enumerate(FEATURE_COLS)
        if x_row[j] == 1
    ]
    return ", ".join(aktif) if aktif else "-"


def build_voting_table(model, X_data, y_data, max_trees=None, classes=None):
    estimators = model.estimators_
    n_trees    = len(estimators)
    n_shown    = min(max_trees, n_trees) if max_trees else n_trees

    print(f"  Jumlah pohon dalam model : {n_trees}")
    print(f"  Pohon yang ditampilkan   : {n_shown}")

    if classes is None:
        classes = model.classes_

    tree_preds_raw = np.array([est.predict(X_data) for est in estimators])

    if tree_preds_raw.dtype.kind in ("i", "u", "f"):
        tree_preds = np.vectorize(lambda idx: classes[int(idx)])(tree_preds_raw)
    else:
        tree_preds = tree_preds_raw

    X_arr  = X_data.values
    y_list = y_data.values
    rows   = []

    for i in range(len(X_data)):
        row = {"ID": i + 1}
        row["Gejala Kerusakan"] = _gejala_aktif(X_arr[i])

        votes = []
        for t in range(n_shown):
            pred = tree_preds[t, i]
            row[f"P{t + 1}"] = _abbrev(pred)
            votes.append(pred)

        for t in range(n_shown, n_trees):
            votes.append(tree_preds[t, i])

        hasil_voting = Counter(votes).most_common(1)[0][0]
        aktual       = y_list[i]
        evaluasi     = "Benar" if hasil_voting == aktual else "Salah"

        row["Hasil Voting"] = _abbrev(hasil_voting)
        row["Aktual"]       = _abbrev(aktual)
        row["Evaluasi"]     = evaluasi
        rows.append(row)

    df_table = pd.DataFrame(rows).set_index("ID")
    return df_table, n_shown


def export_to_excel(df_table, output_path, n_shown):
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_table.to_excel(writer, sheet_name="Voting Mayoritas", index=True)
        ws = writer.sheets["Voting Mayoritas"]

        fill_id   = PatternFill("solid", fgColor="1F4E79")
        fill_feat = PatternFill("solid", fgColor="2E75B6")
        fill_tree = PatternFill("solid", fgColor="375623")
        fill_eval = PatternFill("solid", fgColor="7F6000")
        fill_ok   = PatternFill("solid", fgColor="E2EFDA")
        fill_ng   = PatternFill("solid", fgColor="FCE4D6")

        white_bold = Font(color="FFFFFF", bold=True, size=9)
        small_font = Font(size=8)
        center     = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_wrap  = Alignment(horizontal="left",   vertical="center", wrap_text=True)
        thin = Border(
            left=Side(style="thin"),  right=Side(style="thin"),
            top=Side(style="thin"),   bottom=Side(style="thin"),
        )

        max_col = ws.max_column
        max_row = ws.max_row

        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=1, column=col_idx)
            val  = str(cell.value or "")
            if val == "ID":
                cell.fill = fill_id
            elif val == "Gejala Kerusakan":
                cell.fill = fill_feat
            elif val.startswith("P") and val[1:].isdigit():
                cell.fill = fill_tree
            else:
                cell.fill = fill_eval
            cell.font      = white_bold
            cell.alignment = center
            cell.border    = thin

        # Kolom Evaluasi: 1(ID) + 1(Gejala) + n_shown(pohon) + 3(Hasil,Aktual,Eval)
        eval_col_idx = 1 + 1 + n_shown + 3

        for row_idx in range(2, max_row + 1):
            evaluasi = str(ws.cell(row=row_idx, column=eval_col_idx).value or "")
            row_fill = fill_ok if evaluasi == "Benar" else fill_ng

            for col_idx in range(1, max_col + 1):
                cell        = ws.cell(row=row_idx, column=col_idx)
                cell.fill   = row_fill
                cell.border = thin
                if col_idx == 2:
                    cell.font      = Font(size=8)
                    cell.alignment = left_wrap
                else:
                    cell.font      = small_font
                    cell.alignment = center

        ws.column_dimensions[get_column_letter(1)].width = 5
        ws.column_dimensions[get_column_letter(2)].width = 50
        for c in range(3, 3 + n_shown):
            ws.column_dimensions[get_column_letter(c)].width = 10
        for c in range(3 + n_shown, max_col + 1):
            ws.column_dimensions[get_column_letter(c)].width = 13

        ws.row_dimensions[1].height = 40
        ws.freeze_panes = f"{get_column_letter(3)}2"

    print(f"  File disimpan ke: {output_path}")


def main(max_trees=None, data_split="train"):
    print("=" * 60)
    print("  EXPORT: Tabel Voting Mayoritas per Pohon Keputusan")
    print("=" * 60)

    print(f"\n[1/4] Memuat model dari: {MODEL_PATH}")
    if not os.path.exists(MODEL_PATH):
        print("  ERROR: Jalankan 'python -m app.ml.train' terlebih dahulu.")
        sys.exit(1)
    model          = joblib.load(MODEL_PATH)
    label_encoders = joblib.load(ENCODERS_PATH) if os.path.exists(ENCODERS_PATH) else {}
    print(f"  Model dimuat. Jumlah pohon: {len(model.estimators_)}")

    print(f"\n[2/4] Memuat dataset dari: {DATASET_PATH}")
    df    = load_dataset(DATASET_PATH)
    X_raw = df[FEATURE_COLS].copy()
    y     = df[TARGET_COL].copy()
    X     = transform_features(X_raw.copy(), label_encoders)

    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=0.2, random_state=42, stratify=y
    )

    if data_split == "train":
        X_used, y_used = X_train, y_train
        print(f"  Menggunakan DATA TRAINING: {len(X_used)} baris")
    else:
        X_used, y_used = X_test, y_test
        print(f"  Menggunakan DATA TEST: {len(X_used)} baris")

    print(f"\n[3/4] Membangun tabel voting...")
    df_table, n_shown = build_voting_table(
        model, X_used, y_used,
        max_trees=max_trees,
        classes=model.classes_,
    )

    n_benar = (df_table["Evaluasi"] == "Benar").sum()
    n_salah = (df_table["Evaluasi"] == "Salah").sum()
    acc     = n_benar / len(df_table) * 100

    print(f"\n  Data       : {data_split}")
    print(f"  Akurasi    : {acc:.2f}%")
    print(f"  Benar      : {n_benar}")
    print(f"  Salah      : {n_salah}")

    preview_cols = (
        ["Gejala Kerusakan"]
        + [f"P{i}" for i in range(1, min(4, n_shown + 1))]
        + ["Hasil Voting", "Aktual", "Evaluasi"]
    )
    pd.set_option("display.max_colwidth", 55)
    print(f"\n  Preview (3 baris pertama):")
    print(df_table[preview_cols].head(3).to_string())

    print(f"\n[4/4] Mengekspor ke Excel...")
    export_to_excel(df_table, OUTPUT_PATH, n_shown)
    print("\nSelesai!")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Export tabel voting mayoritas per pohon keputusan."
    )
    parser.add_argument(
        "--max-trees", type=int, default=None, metavar="N",
        help="Jumlah kolom pohon ditampilkan (default: semua).",
    )
    parser.add_argument(
        "--data", choices=["train", "test"], default="train",
        help="Dataset: 'train' (default) atau 'test'.",
    )
    args = parser.parse_args()
    main(max_trees=args.max_trees, data_split=args.data)
