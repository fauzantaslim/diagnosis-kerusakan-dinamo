import io
import csv
import openpyxl
from typing import List, Tuple

from app.models import db
from app.models.dataset import Dataset
from app.ml.random_forest import FEATURE_NAMES, stratified_split


# --------------------------------------------------------------------------- #
#  Helper                                                                      #
# --------------------------------------------------------------------------- #

def _parse_bool(value) -> int:
    """Konversi nilai YA/TIDAK/1/0/True/False ke integer 0 atau 1."""
    if isinstance(value, bool):
        return 1 if value else 0
    if isinstance(value, int):
        return 1 if value else 0
    if isinstance(value, str):
        return 1 if value.strip().upper() in ('YA', '1', 'TRUE', 'YES') else 0
    return 0


def _row_to_kwargs(row: dict) -> dict:
    """
    Memetakan dict baris (dari CSV / XLSX) ke kwargs Dataset.
    Kolom header fleksibel (case-insensitive, strip spasi).
    """
    # Normalisasi key — skip key None (kolom kosong dari Excel)
    norm = {k.strip().lower(): v for k, v in row.items() if k is not None}

    def get(*keys):
        for k in keys:
            if k in norm:
                return norm[k]
        return None

    daya_val        = get('daya (hp)', 'daya', 'daya_hp')
    jumlah_pole_val = get('jumlah pole', 'jumlah_pole', 'pole')
    label_val       = get('label', 'kerusakan', 'diagnosis')

    try:
        daya_val_parsed = float(daya_val) if daya_val not in (None, '', '-') else None
    except (ValueError, TypeError):
        daya_val_parsed = None

    try:
        jumlah_pole_val_parsed = int(jumlah_pole_val) if jumlah_pole_val not in (None, '', '-') else None
    except (ValueError, TypeError):
        jumlah_pole_val_parsed = None

    if not label_val or str(label_val).strip() in ('', '-'):
        raise ValueError("Kolom 'Label' wajib ada dan tidak boleh kosong.")

    return dict(
        daya                                = daya_val_parsed,
        jumlah_pole                         = jumlah_pole_val_parsed,
        suara_bising_abnormal               = _parse_bool(get('suara bising abnormal', 'suara_bising_abnormal')),
        bau_hangus                          = _parse_bool(get('bau hangus', 'bau_hangus')),
        indikasi_overheating                = _parse_bool(get('indikasi overheating', 'indikasi_overheating')),
        putaran_poros_seret                 = _parse_bool(get('putaran poros seret', 'putaran_poros_seret')),
        getaran_berlebih                    = _parse_bool(get('getaran berlebih', 'getaran_berlebih')),
        terminal_overheating                = _parse_bool(get('terminal overheating', 'terminal_overheating')),
        kipas_pendingin_rusak               = _parse_bool(get('kipas pendingin rusak', 'kipas_pendingin_rusak')),
        cooling_duct_tersumbat              = _parse_bool(get('cooling duct tersumbat', 'cooling_duct_tersumbat')),
        resistansi_isolasi_tidak_seimbang   = _parse_bool(get('resistansi isolasi tidak seimbang', 'resistansi_isolasi_tidak_seimbang')),
        resistansi_winding_tidak_seimbang   = _parse_bool(get('resistansi winding tidak seimbang', 'resistansi_winding_tidak_seimbang')),
        arus_antar_fasa_tidak_seimbang      = _parse_bool(get('arus antar fasa tidak seimbang', 'arus_antar_fasa_tidak_seimbang')),
        label                               = str(label_val).strip(),
    )


def _validate_fields(data: dict, require_label: bool = True) -> dict:
    """Validasi dan normalisasi dict dari request JSON."""
    result = {}

    if 'daya' in data:
        result['daya'] = float(data['daya']) if data['daya'] not in (None, '') else None
    if 'jumlah_pole' in data:
        result['jumlah_pole'] = int(data['jumlah_pole']) if data['jumlah_pole'] not in (None, '') else None

    feature_fields = [
        'suara_bising_abnormal', 'bau_hangus', 'indikasi_overheating',
        'putaran_poros_seret', 'getaran_berlebih', 'terminal_overheating',
        'kipas_pendingin_rusak', 'cooling_duct_tersumbat',
        'resistansi_isolasi_tidak_seimbang', 'resistansi_winding_tidak_seimbang',
        'arus_antar_fasa_tidak_seimbang',
    ]
    for field in feature_fields:
        if field in data:
            result[field] = _parse_bool(data[field])

    if 'label' in data:
        result['label'] = str(data['label']).strip()
    elif require_label:
        raise ValueError("Field 'label' wajib diisi.")

    if require_label and 'label' not in result:
        raise ValueError("Field 'label' wajib diisi.")

    return result


# --------------------------------------------------------------------------- #
#  CRUD                                                                        #
# --------------------------------------------------------------------------- #

def get_split_dataset_ids(test_size: float = 0.2, random_state: int = 42) -> Tuple[List[int], List[int]]:
    """Helper untuk mendapatkan list ID dari data training dan testing (single source of truth)."""
    datasets = Dataset.query.order_by(Dataset.id.asc()).all()
    if not datasets:
        return [], []
    
    # Kita hanya butuh ID dan label untuk dimasukkan ke train_test_split
    ids = [d.id for d in datasets]
    labels = [d.label for d in datasets]
    
    from sklearn.model_selection import train_test_split
    train_ids, test_ids, _, _ = train_test_split(ids, labels, test_size=test_size, random_state=random_state, stratify=labels)
    return train_ids, test_ids


def get_all_datasets(page: int = 1, per_page: int = 20, search: str = '', split_filter: str = ''):
    """Ambil semua data dengan pagination dan opsional filter label & split (train/test)."""
    query = Dataset.query

    if search:
        query = query.filter(Dataset.label.ilike(f'%{search}%'))
        
    if split_filter in ('train', 'test'):
        train_ids, test_ids = get_split_dataset_ids()
        if split_filter == 'train':
            query = query.filter(Dataset.id.in_(train_ids))
        elif split_filter == 'test':
            query = query.filter(Dataset.id.in_(test_ids))

    paginated = query.order_by(Dataset.id.asc()).paginate(
        page=page, per_page=per_page, error_out=False
    )

    return {
        'data': [d.to_dict() for d in paginated.items],
        'pagination': {
            'page'    : paginated.page,
            'per_page': paginated.per_page,
            'total'   : paginated.total,
            'pages'   : paginated.pages,
        }
    }


def get_dataset_by_id(dataset_id: int):
    """Ambil satu data berdasarkan id. Return None jika tidak ditemukan."""
    return Dataset.query.get(dataset_id)


def create_dataset(data: dict):
    """Buat record baru. Raise ValueError jika validasi gagal."""
    kwargs = _validate_fields(data, require_label=True)
    dataset = Dataset(**kwargs)
    db.session.add(dataset)
    db.session.commit()
    return dataset


def update_dataset(dataset_id: int, data: dict):
    """Update record. Return None jika tidak ditemukan."""
    dataset = Dataset.query.get(dataset_id)
    if not dataset:
        return None

    kwargs = _validate_fields(data, require_label=False)
    for key, value in kwargs.items():
        setattr(dataset, key, value)

    db.session.commit()
    return dataset


def delete_dataset(dataset_id: int):
    """Hapus record. Return True jika berhasil, False jika tidak ditemukan."""
    dataset = Dataset.query.get(dataset_id)
    if not dataset:
        return False

    db.session.delete(dataset)
    db.session.commit()
    return True


# --------------------------------------------------------------------------- #
#  Data Loading & Split                                                        #
# --------------------------------------------------------------------------- #

def load_split_data(
    test_size: float = 0.2,
    random_state: int = 42,
) -> Tuple[List, List, List, List, List, List]:
    """
    Ambil seluruh dataset dari DB dan lakukan stratified split 80:20.

    Ini adalah **single source of truth** untuk pembagian data training/testing.
    Digunakan oleh ``random_forest_service`` (prediksi) dan
    ``evaluation_service`` (evaluasi GridSearch), sehingga split selalu
    konsisten dengan parameter yang sama.

    Args:
        test_size (float): Proporsi data test. Default ``0.2`` (20 %).
        random_state (int): Seed untuk reproduktibilitas. Default ``42``.

    Returns:
        tuple: ``(X_train, X_test, y_train, y_test, X_all, y_all)``
            - ``X_train`` / ``X_test`` : list of feature vectors (list[list[int]])
            - ``y_train`` / ``y_test`` : list of label strings
            - ``X_all`` / ``y_all``    : seluruh dataset (sebelum split)

    Raises:
        ValueError: Jika jumlah record di DB kurang dari 10.
    """
    datasets = Dataset.query.order_by(Dataset.id.asc()).all()
    if len(datasets) < 10:
        raise ValueError(
            "Dataset terlalu sedikit (minimal 10 data) untuk melakukan split."
        )

    X = [[getattr(d, f) for f in FEATURE_NAMES] for d in datasets]
    y = [d.label for d in datasets]

    from sklearn.model_selection import train_test_split
    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=test_size, random_state=random_state, stratify=y
    )
    return X_train, X_test, y_train, y_test, X, y


# --------------------------------------------------------------------------- #
#  Import CSV / XLSX                                                           #
# --------------------------------------------------------------------------- #

def import_from_file(file) -> dict:
    """
    Parse file CSV atau XLSX kemudian bulk-insert ke tabel datasets.

    Returns:
        dict: { 'imported': int, 'skipped': int, 'errors': list }
    """
    filename = file.filename.lower()
    rows = []

    if filename.endswith('.csv'):
        content = file.read().decode('utf-8-sig')  # handle BOM
        reader = csv.DictReader(io.StringIO(content))
        rows = list(reader)

    elif filename.endswith(('.xlsx', '.xls')):

        wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
        ws = wb.active

        headers = [cell.value for cell in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            # Lewati baris kosong
            if all(v is None for v in row):
                continue
            rows.append(dict(zip(headers, row)))
    else:
        raise ValueError("Format file tidak didukung. Gunakan .csv atau .xlsx")

    imported = 0
    skipped  = 0
    errors   = []

    for idx, row in enumerate(rows, start=2):  # start=2 karena baris 1 = header
        try:
            kwargs  = _row_to_kwargs(row)
            dataset = Dataset(**kwargs)
            db.session.add(dataset)
            imported += 1
        except Exception as e:
            skipped += 1
            errors.append({'row': idx, 'reason': str(e)})

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        raise RuntimeError(f"Gagal menyimpan ke database: {str(e)}")

    return {'imported': imported, 'skipped': skipped, 'errors': errors}
